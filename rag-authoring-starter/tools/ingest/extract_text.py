#!/usr/bin/env python3
"""
Extract plain text from the corpus sources listed in rag/sources.yaml.

Supports DOCX (via python-docx) and PDF (via pypdf). Each source is written to
rag/corpus/<...>.extracted.txt alongside a small JSON sidecar carrying
provenance metadata (id, title, status, target_part, source path).

The script degrades gracefully:
- If a source file is missing, it is skipped and reported.
- If an optional dependency is missing, that format is skipped and reported.

Usage:
    python tools/ingest/extract_text.py            # extract everything present
    python tools/ingest/extract_text.py --id ...   # extract a single source id
"""

from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parents[2]
SOURCES = ROOT / "rag" / "sources.yaml"


def load_sources() -> list[dict[str, Any]]:
    text = SOURCES.read_text(encoding="utf-8")
    try:
        import yaml  # type: ignore

        data = yaml.safe_load(text) or {}
        return list(data.get("sources", []))
    except Exception:
        # Reuse the fallback parser from build_inventory for consistency.
        from build_inventory import _fallback_yaml_sources  # type: ignore

        return _fallback_yaml_sources(text)


def extract_docx(path: Path) -> str:
    try:
        import docx  # type: ignore
    except Exception as exc:  # pragma: no cover - dependency guidance
        raise RuntimeError(
            "python-docx is required for DOCX extraction. "
            "Install with: pip install python-docx"
        ) from exc

    document = docx.Document(str(path))
    parts: list[str] = []
    for para in document.paragraphs:
        style = (para.style.name or "").lower() if para.style else ""
        text = para.text.strip()
        if not text:
            continue
        if style.startswith("heading"):
            parts.append("\n" + text)
        else:
            parts.append(text)
    # Include table text, which DASH-IF specs use heavily.
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def extract_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader  # type: ignore
    except Exception as exc:  # pragma: no cover - dependency guidance
        raise RuntimeError(
            "pypdf is required for PDF extraction. Install with: pip install pypdf"
        ) from exc

    reader = PdfReader(str(path))
    parts: list[str] = []
    for page in reader.pages:
        parts.append(page.extract_text() or "")
    return "\n".join(parts)


def extract_one(source: dict[str, Any]) -> tuple[str, str]:
    """Return (status, message) for a single source."""
    corpus_path = str(source.get("corpus_path", "")).strip()
    if not corpus_path:
        return "skip", "no corpus_path"

    src = ROOT / corpus_path
    if not src.exists():
        return "missing", f"file not found: {corpus_path}"

    suffix = src.suffix.lower()
    try:
        if suffix == ".docx":
            text = extract_docx(src)
        elif suffix == ".pdf":
            text = extract_pdf(src)
        elif suffix in {".txt", ".md"}:
            text = src.read_text(encoding="utf-8", errors="replace")
        else:
            return "skip", f"unsupported format: {suffix}"
    except RuntimeError as exc:
        return "error", str(exc)

    out_txt = src.with_suffix(src.suffix + ".extracted.txt")
    out_txt.write_text(text, encoding="utf-8")

    sidecar = {
        "id": source.get("id", ""),
        "title": source.get("title", ""),
        "type": source.get("type", ""),
        "status": source.get("status", ""),
        "target_part": source.get("target_part", ""),
        "source_path": corpus_path,
        "extracted_path": str(out_txt.relative_to(ROOT)).replace("\\", "/"),
        "char_count": len(text),
        "tags": source.get("tags", []),
    }
    out_txt.with_suffix(".json").write_text(
        json.dumps(sidecar, indent=2), encoding="utf-8"
    )
    return "ok", f"{len(text)} chars -> {out_txt.relative_to(ROOT)}"


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--id", help="Extract only the source with this id")
    args = parser.parse_args()

    sources = load_sources()
    if args.id:
        sources = [s for s in sources if s.get("id") == args.id]
        if not sources:
            print(f"No source with id={args.id!r}")
            return

    counts = {"ok": 0, "missing": 0, "skip": 0, "error": 0}
    for source in sources:
        status, message = extract_one(source)
        counts[status] = counts.get(status, 0) + 1
        print(f"[{status:>7}] {source.get('id', '?'):<32} {message}")

    print(
        "\nSummary: "
        + ", ".join(f"{k}={v}" for k, v in counts.items() if v)
        + f" (of {len(sources)} sources)"
    )


if __name__ == "__main__":
    main()
